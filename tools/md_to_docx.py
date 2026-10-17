"""
Markdown → Word(.docx) 转换器 —— 让论文材料能从纯文本重新生成。

为什么需要这个脚本:
    开题报告书、文献综述、论文正文这些材料要交给老师和教务系统,
    必须是可以直接打开的 .docx;但 .docx 是二进制,改一个字看不出改了哪里,
    也没法 diff。所以约定:论文材料一律先用 Markdown 写(可版本管理、
    可 diff、可 grep),再一条命令导出 Word。
    这样"论文里的每个数字都能重跑出来"这条原则才延伸到文档层。

用法:
    .python\\python.exe tools\\md_to_docx.py 论文\\开题报告书.md
    .python\\python.exe tools\\md_to_docx.py 论文\\开题报告书.md -o 论文\\开题报告书.docx

支持的语法(够用即可,不做完整 Markdown 实现):
    # / ## / ### / ####     标题(# 为文档标题)
    **粗体**  `等宽`         行内格式
    - / * / +                无序列表
    1. 2.                    有序列表
    >                        引用
    | a | b |                表格(第二行为 ---|--- 分隔行)
    ```                      代码块(等宽字体,保留缩进)
    ![说明](路径)            插图,独占一段并居中;路径相对 .md 文件所在目录
    ---                      分隔线(忽略)

退出码 0 = 成功。
"""

import argparse
import os
import re
import sys

_HERE = os.path.dirname(os.path.abspath(__file__))
_WS = os.path.dirname(_HERE)

# 中文字体必须显式设置 w:eastAsia,只设 run.font.name 只影响 ASCII,
# 中文会退回默认字体 —— 这是 python-docx 最常见的坑。
BODY_FONT = '宋体'
HEAD_FONT = '黑体'
MONO_FONT = 'Consolas'

BODY_SIZE = 12.0     # 小四
CODE_SIZE = 10.5     # 五号

# 插图宽度上限:论文正文版心约 14.6cm(21 - 3.17*2),留一点余量。
MAX_IMAGE_WIDTH_CM = 14.0

_RE_INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')
_RE_ULIST = re.compile(r'^\s*[-*+]\s+(.*)$')
_RE_OLIST = re.compile(r'^\s*(\d+)[.)]\s+(.*)$')
_RE_HEAD = re.compile(r'^(#{1,6})\s+(.*)$')
_RE_TABLE_SEP = re.compile(r'^\s*\|?[\s:|-]+\|[\s:|-]*$')
_RE_IMAGE = re.compile(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$')


# ------------------------------------------------------------------ 字体工具

def set_font(run, name=BODY_FONT, size=None, bold=None, mono=False):
    """设置 run 的字体,中英文都覆盖。

    中文必须通过 w:eastAsia 指定,否则 Word 用默认中文字体渲染。
    """
    from docx.oxml.ns import qn
    from docx.shared import Pt

    ascii_name = MONO_FONT if mono else name
    run.font.name = ascii_name
    rpr = run._element.get_or_add_rPr()
    # get_or_add_rFonts() 会在缺失时新建,直接 .rFonts 可能是 None
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn('w:ascii'), ascii_name)
    rfonts.set(qn('w:hAnsi'), ascii_name)
    # 代码块里的制表符/框线字符走东亚字体,否则会显示成方框
    rfonts.set(qn('w:eastAsia'), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def add_runs(paragraph, text, size=None):
    """把一行文本写进段落,识别 **粗体** 与 `等宽`。"""
    for part in _RE_INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, bold=True)
        elif part.startswith('`') and part.endswith('`') and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=size, mono=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size)


# ------------------------------------------------------------------ 表格解析

def split_row(line):
    """把 `| a | b |` 切成 ['a', 'b']。"""
    s = line.strip()
    if s.startswith('|'):
        s = s[1:]
    if s.endswith('|'):
        s = s[:-1]
    return [c.strip() for c in s.split('|')]


def is_table_sep(line):
    """判断是不是表格的分隔行(|---|---|)。"""
    s = line.strip()
    if '|' not in s:
        return False
    return bool(_RE_TABLE_SEP.match(s))


def resolve_image_path(raw_path, base_dir):
    """把 Markdown 里的图片路径解析成绝对路径。

    支持相对 .md 所在目录、相对工作区根目录、以及绝对路径三种写法。
    找不到就返回 None,由调用方决定报错还是跳过。
    """
    path = raw_path.strip().strip('<>')
    if os.path.isabs(path):
        return path if os.path.isfile(path) else None
    for root in (base_dir, _WS):
        candidate = os.path.join(root, path)
        if os.path.isfile(candidate):
            return candidate
    return None


def image_size_cm(path, max_width=MAX_IMAGE_WIDTH_CM):
    """按版心宽度等比缩放,返回 (宽, 高),单位 cm。

    只限宽不限高会出事:一张竖长的架构图按 14cm 宽缩放后高度可能超过一页,
    Word 会把图截断或挤到下一页留半页空白。所以高度也要封顶。
    """
    from PIL import Image

    with Image.open(path) as im:
        px_w, px_h = im.size
    if px_w <= 0 or px_h <= 0:
        return max_width, max_width

    # 按 96 DPI 估算原始物理尺寸,再决定是否缩放
    width_cm = px_w / 96.0 * 2.54
    height_cm = px_h / 96.0 * 2.54

    scale = 1.0
    if width_cm > max_width:
        scale = min(scale, max_width / width_cm)
    max_height = 20.0   # 版心高度约 24.6cm,留出图注空间
    if height_cm * scale > max_height:
        scale = min(scale, max_height / height_cm)
    return width_cm * scale, height_cm * scale


# ------------------------------------------------------------------ 主转换

def convert(md_text, out_path, base_dir=None):
    """把 Markdown 文本写成 .docx,返回写入路径。

    base_dir 用于解析 ![说明](相对路径) 里的图片路径;
    默认取当前工作目录。
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt

    if base_dir is None:
        base_dir = os.getcwd()

    doc = Document()

    # 页面: A4 + 论文常用页边距
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.17)
        section.right_margin = Cm(3.17)

    # 正文样式: 宋体小四、1.5 倍行距
    normal = doc.styles['Normal']
    normal.font.size = Pt(BODY_SIZE)
    normal.paragraph_format.line_spacing = 1.5
    # 样式里的 rPr 可能没有 rFonts 子元素,必须用 get_or_add_rFonts()
    from docx.oxml.ns import qn
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(
        qn('w:eastAsia'), BODY_FONT)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        # --- 代码块
        if stripped.startswith('```'):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1  # 跳过结束的 ```
            for cl in buf:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(cl if cl.strip() else ' ')
                set_font(run, size=CODE_SIZE, mono=True)
            continue

        # --- 插图
        m = _RE_IMAGE.match(stripped)
        if m:
            alt, raw_path = m.group(1), m.group(2)
            found = resolve_image_path(raw_path, base_dir)
            if found is None:
                # 不静默跳过:漏图是定稿前必须发现的问题,
                # 在 docx 里留一个显眼的标记比悄悄少一张图好。
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run('[缺图: %s]' % raw_path)
                set_font(run, size=BODY_SIZE, bold=True)
                i += 1
                continue
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            width_cm, _ = image_size_cm(found)
            p.add_run().add_picture(found, width=Cm(width_cm))
            if alt:
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = cap.add_run(alt)
                set_font(run, size=BODY_SIZE - 1)
            i += 1
            continue

        # --- 表格
        if stripped.startswith('|') and i + 1 < len(lines) \
                and is_table_sep(lines[i + 1]):
            header = split_row(lines[i])
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            ncol = len(header)
            table = doc.add_table(rows=1 + len(rows), cols=ncol)
            table.style = 'Table Grid'
            for c, text in enumerate(header):
                cell = table.cell(0, c)
                cell.text = ''
                add_runs(cell.paragraphs[0], text, size=BODY_SIZE - 1)
                for run in cell.paragraphs[0].runs:
                    run.font.bold = True
            for r, row in enumerate(rows, start=1):
                for c in range(ncol):
                    cell = table.cell(r, c)
                    cell.text = ''
                    val = row[c] if c < len(row) else ''
                    add_runs(cell.paragraphs[0], val, size=BODY_SIZE - 1)
            doc.add_paragraph()
            continue

        # --- 标题
        m = _RE_HEAD.match(stripped)
        if m:
            level = len(m.group(1))
            style = 'Title' if level == 1 else 'Heading %d' % (level - 1)
            p = doc.add_paragraph(style=style)
            add_runs(p, m.group(2))
            for run in p.runs:
                set_font(run, name=HEAD_FONT)
            i += 1
            continue

        # --- 无序列表
        m = _RE_ULIST.match(line)
        if m:
            p = doc.add_paragraph(style='List Bullet')
            add_runs(p, m.group(1))
            i += 1
            continue

        # --- 有序列表
        m = _RE_OLIST.match(line)
        if m:
            p = doc.add_paragraph(style='List Number')
            add_runs(p, m.group(2))
            i += 1
            continue

        # --- 引用
        if stripped.startswith('>'):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.74)
            add_runs(p, stripped.lstrip('> ').strip())
            for run in p.runs:
                run.font.italic = True
            i += 1
            continue

        # --- 分隔线 / 空行
        if not stripped or set(stripped) <= set('-*_') and len(stripped) >= 3:
            i += 1
            continue

        # --- 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(BODY_SIZE * 2)
        add_runs(p, stripped)
        i += 1

    out_dir = os.path.dirname(os.path.abspath(out_path))
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    doc.save(out_path)
    return out_path


def main(argv=None):
    ap = argparse.ArgumentParser(description='Markdown → Word(.docx)')
    ap.add_argument('input', help='输入的 .md 文件')
    ap.add_argument('-o', '--out', default=None,
                    help='输出的 .docx(默认与输入同目录同名)')
    args = ap.parse_args(argv)

    if not os.path.isfile(args.input):
        raise SystemExit('输入文件不存在: %s' % args.input)

    with open(args.input, encoding='utf-8') as fh:
        md_text = fh.read()

    out = args.out
    if not out:
        base = os.path.splitext(args.input)[0]
        out = base + '.docx'

    # 图片路径相对 .md 所在目录解析,这样 论文/毕业论文初稿.md 里的
    # figures/fig-3-1.png 能直接找到,不用写成 论文/figures/...
    convert(md_text, out, base_dir=os.path.dirname(os.path.abspath(args.input)))
    print('已生成 %s' % os.path.relpath(out, _WS))
    return 0


if __name__ == '__main__':
    sys.exit(main())
