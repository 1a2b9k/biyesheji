"""
Markdown → Word 转换器的测试(tools/md_to_docx.py)。

这个工具产出的是要交给老师和教务系统的正式材料,格式错了很难发现
(Word 打开才看得出来),所以把几条最容易错的约定锁死在测试里:

  1. **中文字体必须写 w:eastAsia。** 只设 run.font.name 只影响 ASCII,
     中文会退回默认字体 —— 这是 python-docx 最经典的坑。
  2. **表格分隔行不能被当成正文。** `|---|---|` 若当普通段落渲染,
     论文表格里会多出一行乱码似的竖线。
  3. **代码块要原样保留缩进。** 技术路线图靠缩进表达层级,被吃掉就废了。
  4. **行内 `**粗体**` 与 `` `等宽` `` 要拆成多个 run**,不能被当作字面文本。
  5. **表格列数不齐时要补齐而不是崩。** 手写 Markdown 表格很容易少写一列。

不加载 Word、不做渲染:只读回 .docx 的 XML 结构做断言。
"""

import os
import re
import sys
import zipfile

import pytest

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(
    os.path.abspath(__file__))))
_TOOLS = os.path.join(_REPO_ROOT, 'tools')
if _TOOLS not in sys.path:
    sys.path.insert(0, _TOOLS)

import md_to_docx  # noqa: E402

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


# ---------------------------------------------------------------- 辅助

def _convert(tmp_path, md_text, name='out.docx'):
    out = str(tmp_path / name)
    md_to_docx.convert(md_text, out)
    return out


def _document_xml(path):
    with zipfile.ZipFile(path) as z:
        return z.read('word/document.xml').decode('utf-8')


def _paragraph_texts(path):
    """返回文档里每个段落的纯文本。"""
    from docx import Document
    doc = Document(path)
    return [p.text for p in doc.paragraphs]


def _east_asia_fonts(path):
    """收集文档里所有出现过的 w:eastAsia 取值。"""
    import re
    xml = _document_xml(path)
    return set(re.findall(r'w:eastAsia="([^"]+)"', xml))


# ---------------------------------------------------------------- 行内解析

def test_split_row_strips_outer_pipes():
    assert md_to_docx.split_row('| a | b |') == ['a', 'b']
    assert md_to_docx.split_row('a | b') == ['a', 'b']
    assert md_to_docx.split_row('| 只有一个 |') == ['只有一个']


def test_split_row_keeps_empty_cells():
    assert md_to_docx.split_row('| a |  | c |') == ['a', '', 'c']


@pytest.mark.parametrize('line,expected', [
    ('|---|---|', True),
    ('| --- | :---: |', True),
    ('|:--|--:|', True),
    ('| 内容 | 内容 |', False),
    ('没有竖线', False),
])
def test_is_table_sep(line, expected):
    assert md_to_docx.is_table_sep(line) is expected


def test_inline_regex_splits_bold_and_code():
    parts = md_to_docx._RE_INLINE.split('普通 **粗体** 与 `代码` 结束')
    assert '**粗体**' in parts
    assert '`代码`' in parts
    assert '普通 ' in parts


# ---------------------------------------------------------------- 字体

def test_chinese_font_is_written_as_east_asia(tmp_path):
    """中文必须落到 w:eastAsia,否则 Word 用默认字体渲染。"""
    path = _convert(tmp_path, '# 标题\n\n这是一段中文正文。\n')
    fonts = _east_asia_fonts(path)
    assert md_to_docx.BODY_FONT in fonts
    assert md_to_docx.HEAD_FONT in fonts


def test_code_block_uses_mono_font(tmp_path):
    path = _convert(tmp_path, '```\n数据 → 模型\n```\n')
    xml = _document_xml(path)
    assert md_to_docx.MONO_FONT in xml


def test_normal_style_has_east_asia_font(tmp_path):
    """Normal 样式的 rPr 里若没有 rFonts,中文会退回默认字体。"""
    path = _convert(tmp_path, '正文\n')
    xml = _document_xml(path)
    # 样式定义在 styles.xml
    with zipfile.ZipFile(path) as z:
        styles = z.read('word/styles.xml').decode('utf-8')
    assert 'w:eastAsia="%s"' % md_to_docx.BODY_FONT in styles


# ---------------------------------------------------------------- 结构

def test_headings_map_to_levels(tmp_path):
    path = _convert(tmp_path, '# 文档标题\n\n## 一级\n\n### 二级\n')
    from docx import Document
    doc = Document(path)
    styles = [p.style.name for p in doc.paragraphs if p.text.strip()]
    assert styles[0] == 'Title'
    assert 'Heading 1' in styles
    assert 'Heading 2' in styles


def test_table_becomes_real_table(tmp_path):
    md = '| 任务 | 准确率 |\n|---|---|\n| A | 0.754 |\n| B | 0.843 |\n'
    path = _convert(tmp_path, md)
    from docx import Document
    doc = Document(path)
    assert len(doc.tables) == 1
    t = doc.tables[0]
    assert len(t.rows) == 3
    assert len(t.columns) == 2
    assert t.cell(0, 0).text == '任务'
    assert t.cell(2, 1).text == '0.843'


def test_table_separator_not_rendered_as_text(tmp_path):
    """分隔行 |---|---| 不能变成正文段落。"""
    md = '| a | b |\n|---|---|\n| 1 | 2 |\n'
    path = _convert(tmp_path, md)
    texts = _paragraph_texts(path)
    assert not any('---' in t for t in texts)


def test_table_ragged_rows_are_padded(tmp_path):
    """行内列数少于表头时补齐,不能崩。"""
    md = '| a | b | c |\n|---|---|---|\n| 1 |\n'
    path = _convert(tmp_path, md)
    from docx import Document
    doc = Document(path)
    t = doc.tables[0]
    assert len(t.columns) == 3
    assert t.cell(1, 1).text == ''


def test_ragged_row_with_extra_cells_is_ignored(tmp_path):
    """行内列数多于表头时,多余的丢弃,不能崩。"""
    md = '| a | b |\n|---|---|\n| 1 | 2 | 3 | 4 |\n'
    path = _convert(tmp_path, md)
    from docx import Document
    doc = Document(path)
    assert len(doc.tables[0].columns) == 2


def test_code_block_preserves_indentation(tmp_path):
    """技术路线图靠缩进表达层级。"""
    md = '```\n数据获取\n        │\n        ├─→ 路线一\n```\n'
    path = _convert(tmp_path, md)
    texts = _paragraph_texts(path)
    joined = '\n'.join(texts)
    assert '        ├─→ 路线一' in joined


def test_lists_become_list_styles(tmp_path):
    md = '- 第一项\n- 第二项\n\n1. 有序一\n2. 有序二\n'
    path = _convert(tmp_path, md)
    from docx import Document
    doc = Document(path)
    styles = [p.style.name for p in doc.paragraphs if p.text.strip()]
    assert 'List Bullet' in styles
    assert 'List Number' in styles


def test_ordered_list_does_not_keep_the_number(tmp_path):
    """编号交给 Word 的列表样式,文本里不能残留 '1.'。"""
    path = _convert(tmp_path, '1. 第一条\n')
    texts = [t for t in _paragraph_texts(path) if t.strip()]
    assert '第一条' in texts
    assert not any(t.strip().startswith('1.') for t in texts)


def test_bold_run_is_bold(tmp_path):
    path = _convert(tmp_path, '普通 **重点** 收尾\n')
    from docx import Document
    doc = Document(path)
    para = [p for p in doc.paragraphs if '重点' in p.text][0]
    bold_runs = [r for r in para.runs if r.text == '重点']
    assert bold_runs and bold_runs[0].font.bold is True


def test_inline_code_run_uses_mono(tmp_path):
    path = _convert(tmp_path, '切换 `FIN_SENTIMENT_MODEL` 环境变量\n')
    xml = _document_xml(path)
    assert md_to_docx.MONO_FONT in xml


def test_blockquote_is_indented(tmp_path):
    path = _convert(tmp_path, '> 这是一条提示\n')
    from docx import Document
    from docx.shared import Cm
    doc = Document(path)
    para = [p for p in doc.paragraphs if '提示' in p.text][0]
    # Word 把长度存成 twip,再读回来与 Cm() 差几百 EMU 是正常的
    assert abs(para.paragraph_format.left_indent - Cm(0.74)) < 1000


def test_horizontal_rule_is_dropped(tmp_path):
    path = _convert(tmp_path, '上\n\n---\n\n下\n')
    texts = _paragraph_texts(path)
    assert '上' in texts and '下' in texts
    assert not any(set(t.strip()) <= set('-') and t.strip() for t in texts)


def test_empty_document_does_not_crash(tmp_path):
    path = _convert(tmp_path, '')
    assert os.path.isfile(path)


def test_page_is_a4_with_thesis_margins(tmp_path):
    path = _convert(tmp_path, '正文\n')
    from docx import Document
    from docx.shared import Cm
    s = Document(path).sections[0]
    # Word 以 twip 存储长度,读回来与 Cm() 有几百 EMU 的舍入差
    assert abs(s.page_width - Cm(21.0)) < 1000
    assert abs(s.page_height - Cm(29.7)) < 1000
    assert abs(s.left_margin - Cm(3.17)) < 1000


# ---------------------------------------------------------------- 插图

def _make_png(path, size=(800, 400)):
    """造一张纯色 PNG,尺寸可控,用来验证插图与缩放。"""
    from PIL import Image
    Image.new('RGB', size, (200, 180, 150)).save(str(path))
    return str(path)


def test_image_is_embedded_into_docx(tmp_path):
    img = _make_png(tmp_path / 'a.png')
    md = '上文\n\n![图 1-1 测试图](%s)\n\n下文\n' % img
    path = _convert(tmp_path, md)
    xml = _document_xml(path)
    # 真插图会在 document.xml 里留下 drawing 与图片关系引用
    assert '<w:drawing>' in xml or '<w:pict>' in xml
    assert 'r:embed' in xml


def test_image_caption_is_the_alt_text(tmp_path):
    img = _make_png(tmp_path / 'a.png')
    md = '![图 3-1 训练集标签分布](%s)\n' % img
    path = _convert(tmp_path, md)
    texts = _paragraph_texts(path)
    assert '图 3-1 训练集标签分布' in texts


def test_image_without_alt_has_no_caption(tmp_path):
    img = _make_png(tmp_path / 'a.png')
    path = _convert(tmp_path, '![]( %s )\n' % img)
    texts = _paragraph_texts(path)
    # 只该有图片本身,不该多出一个空图注段落
    assert all(t.strip() == '' or '[缺图' in t for t in texts)


def test_missing_image_leaves_visible_marker(tmp_path):
    path = _convert(tmp_path, '![图 9-9 不存在](nope/zzz.png)\n')
    texts = _paragraph_texts(path)
    joined = '|'.join(texts)
    # 漏图必须显眼,不能被悄悄跳过
    assert '[缺图' in joined and 'nope/zzz.png' in joined


def test_resolve_image_path_tries_base_dir_then_workspace(tmp_path):
    sub = tmp_path / 'sub'
    sub.mkdir()
    img = _make_png(sub / 'x.png')
    # 相对 base_dir 能解析
    got = md_to_docx.resolve_image_path('x.png', str(sub))
    assert got and os.path.samefile(got, img)
    # 不存在则返回 None,由调用方决定怎么处理
    assert md_to_docx.resolve_image_path('nope.png', str(sub)) is None
    # 绝对路径直接命中
    assert md_to_docx.resolve_image_path(img, str(tmp_path)) == img


def test_image_size_caps_width(tmp_path):
    img = _make_png(tmp_path / 'wide.png', size=(4000, 500))
    width_cm, height_cm = md_to_docx.image_size_cm(img)
    assert width_cm <= md_to_docx.MAX_IMAGE_WIDTH_CM + 1e-6
    # 等比缩放,宽高比保持
    assert abs(width_cm / height_cm - 8.0) < 0.05


def test_image_size_caps_height_for_tall_images(tmp_path):
    # 只限宽会让竖长图超过一页被 Word 截断,所以高度也要封顶
    img = _make_png(tmp_path / 'tall.png', size=(300, 6000))
    width_cm, height_cm = md_to_docx.image_size_cm(img)
    assert height_cm <= 20.0 + 1e-6
    assert width_cm <= md_to_docx.MAX_IMAGE_WIDTH_CM + 1e-6


def test_image_size_keeps_small_image_unscaled(tmp_path):
    img = _make_png(tmp_path / 'small.png', size=(96, 96))
    width_cm, height_cm = md_to_docx.image_size_cm(img)
    # 96px @96DPI = 1 英寸 = 2.54cm,小图不该被放大
    assert abs(width_cm - 2.54) < 0.01
    assert abs(height_cm - 2.54) < 0.01


def test_convert_resolves_image_relative_to_md_directory(tmp_path):
    # 论文里写 figures/fig-3-1.png,相对 .md 所在目录解析
    doc_dir = tmp_path / 'paper'
    (doc_dir / 'figures').mkdir(parents=True)
    _make_png(doc_dir / 'figures' / 'f.png')
    md = '![图 1-1 相对路径](figures/f.png)\n'
    out = str(tmp_path / 'out.docx')
    md_to_docx.convert(md, out, base_dir=str(doc_dir))
    assert '[缺图' not in '|'.join(_paragraph_texts(out))


# ---------------------------------------------------------------- 目录域

_TOC_SAMPLE = """# 论文题目

<!-- TOC -->

## 第 1 章 绪论

正文。

### 1.1 研究背景

小节正文。

## 第 2 章 相关技术

结尾。
"""


def _settings_xml(path):
    with zipfile.ZipFile(path) as z:
        return z.read('word/settings.xml').decode('utf-8')


def test_toc_marker_becomes_a_real_field(tmp_path):
    """目录必须是 Word 的 TOC 域,而不是拼好的死文字。

    死文字在标题改动后就错,而且错得不明显 —— 只有域才能在 Word 里按 F9
    同步。这里查 instrText 里的域指令是否写进去了。
    """
    out = _convert(tmp_path, _TOC_SAMPLE)
    xml = _document_xml(out)
    assert 'TOC \\o "1-3"' in xml, '没有写入 TOC 域指令'


def test_toc_field_has_begin_separate_and_end(tmp_path):
    """域的 fldChar 必须三段齐全。

    缺 separate 时 Word 认为域还没开始计算结果,打开文档会报错 ——
    这是手写域最容易漏的一段。
    """
    out = _convert(tmp_path, _TOC_SAMPLE)
    xml = _document_xml(out)
    assert xml.count('w:fldCharType="begin"') == 1
    assert xml.count('w:fldCharType="separate"') == 1
    assert xml.count('w:fldCharType="end"') == 1


def test_toc_field_has_placeholder_text(tmp_path):
    """域结果留占位提示,不能是空白。

    空白会让人以为目录坏了;一句提示至少说明了该按 F9。
    """
    out = _convert(tmp_path, _TOC_SAMPLE)
    text = '|'.join(_paragraph_texts(out))
    assert 'F9' in text


def test_toc_title_is_not_a_heading(tmp_path):
    """「目录」二字不能出现在它自己的目录里。

    目录标题若用 Heading 样式,就会被 TOC 域收成一级条目,
    出现「目录 … 目录」这种自指。所以它是普通段落 + 手工字体。
    """
    out = _convert(tmp_path, _TOC_SAMPLE)
    xml = _document_xml(out)
    # 找到「目　录」所在的段落,确认它没有 pStyle=Heading
    for para in re.findall(r'<w:p\b.*?</w:p>', xml, re.S):
        if '目　录' in para:
            assert 'Heading' not in para and 'w:pStyle w:val="Title"' not in para, \
                '目录标题用了标题样式,会自指'
            break
    else:
        raise AssertionError('没有找到目录标题')


def test_update_fields_is_enabled_when_toc_present(tmp_path):
    """有目录时必须让 Word 打开就更新域,否则看到的只是占位提示。"""
    out = _convert(tmp_path, _TOC_SAMPLE)
    assert 'updateFields' in _settings_xml(out)


def test_update_fields_is_absent_without_toc(tmp_path):
    """没有目录的文档不该带上这个设置,否则 Word 会多问一次「是否更新」。"""
    out = _convert(tmp_path, '# 只有标题\n\n正文\n')
    assert 'updateFields' not in _settings_xml(out)


def test_toc_marker_is_not_rendered_as_text(tmp_path):
    """`<!-- TOC -->` 是给转换器的指令,不该原样出现在正文里。"""
    out = _convert(tmp_path, _TOC_SAMPLE)
    text = '|'.join(_paragraph_texts(out))
    assert '<!-- TOC -->' not in text


def test_toc_marker_only_takes_effect_once(tmp_path):
    """同一文档里写了两个标记也只插一个目录。"""
    md = '# 题目\n\n<!-- TOC -->\n\n## 第一章\n\n正文\n\n<!-- TOC -->\n\n## 第二章\n'
    out = _convert(tmp_path, md)
    xml = _document_xml(out)
    assert xml.count('w:fldCharType="begin"') == 1


def test_toc_ignores_case_and_spacing(tmp_path):
    """`<!--toc-->` 这种写法也应识别,免得手写时大小写不一致就静默失效。"""
    md = '# 题目\n\n<!--toc-->\n\n## 第一章\n\n正文\n'
    out = _convert(tmp_path, md)
    assert 'TOC \\o "1-3"' in _document_xml(out)


# ---------------------------------------------------------------- CLI

def test_main_writes_default_output(tmp_path, monkeypatch, capsys):
    src = tmp_path / 'x.md'
    src.write_text('# 标题\n\n正文\n', encoding='utf-8')
    monkeypatch.chdir(tmp_path)
    assert md_to_docx.main([str(src)]) == 0
    assert (tmp_path / 'x.docx').is_file()


def test_main_respects_out_flag(tmp_path):
    src = tmp_path / 'x.md'
    src.write_text('正文\n', encoding='utf-8')
    dst = tmp_path / 'sub' / 'y.docx'
    assert md_to_docx.main([str(src), '-o', str(dst)]) == 0
    assert dst.is_file()


def test_main_missing_input_exits(tmp_path):
    with pytest.raises(SystemExit):
        md_to_docx.main([str(tmp_path / 'nope.md')])


# ---------------------------------------------------------------- 真实材料

def test_real_proposal_material_converts(tmp_path):
    """仓库里的开题报告书必须能一键转成 Word。"""
    src = os.path.join(_REPO_ROOT, '论文', '开题报告书.md')
    if not os.path.isfile(src):
        pytest.skip('尚无 论文/开题报告书.md')
    with open(src, encoding='utf-8') as fh:
        text = fh.read()
    out = _convert(tmp_path, text, 'proposal.docx')
    from docx import Document
    doc = Document(out)
    # 必须含真实实验数字,防止手抄错
    joined = '\n'.join(p.text for p in doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            joined += '\n' + ' '.join(c.text for c in row.cells)
    assert '0.7540' in joined
    assert '8516' in joined


# ---------------------------------------------------------------- 引用完整性

def _proposal_text():
    src = os.path.join(_REPO_ROOT, '论文', '开题报告书.md')
    if not os.path.isfile(src):
        pytest.skip('尚无 论文/开题报告书.md')
    with open(src, encoding='utf-8') as fh:
        return fh.read()


def _split_refs(text):
    """把正文与参考文献表切开,返回 (正文, {编号: 条目行})。"""
    marker = '## 七、主要参考文献'
    assert marker in text, '开题报告书缺少参考文献章节'
    body, _, refs = text.partition(marker)
    entries = {}
    for line in refs.splitlines():
        m = re.match(r'^\[(\d+)\]\s+(.*)$', line.strip())
        if m:
            entries[int(m.group(1))] = m.group(2)
    return body, entries


def test_every_cited_number_has_a_reference_entry():
    """正文里引用的 [n] 必须在文末参考文献表里有对应条目。

    漏一条就是"引而不注",答辩时被问住很难解释。
    """
    body, entries = _split_refs(_proposal_text())
    cited = {int(n) for n in re.findall(r'\[(\d+)\]', body)}
    assert cited, '正文里一处引用都没有,文献综述还没写'
    missing = sorted(cited - set(entries))
    assert not missing, f'正文引用了但文末没有的编号: {missing}'


def test_no_orphan_reference_entries():
    """参考文献表里不能有正文从未引用的条目(凑数嫌疑)。"""
    body, entries = _split_refs(_proposal_text())
    cited = {int(n) for n in re.findall(r'\[(\d+)\]', body)}
    orphan = sorted(set(entries) - cited)
    assert not orphan, f'列在文末但正文从未引用的编号: {orphan}'


def test_reference_numbers_are_contiguous():
    """编号必须从 1 开始连续,不能跳号 —— 跳号说明删条目时漏改了。"""
    _, entries = _split_refs(_proposal_text())
    assert entries, '参考文献表是空的'
    nums = sorted(entries)
    assert nums == list(range(1, len(nums) + 1)), f'编号不连续: {nums}'


def test_reference_entries_look_like_gbt7714():
    """每条至少要有作者、题名与出处三要素,不能只写一个标题。"""
    _, entries = _split_refs(_proposal_text())
    for num, entry in sorted(entries.items()):
        assert len(entry) > 25, f'[{num}] 著录过短,缺出处信息: {entry}'
        # GB/T 7714 用 [J]/[C]/[EB/OL] 等文献类型标识
        assert re.search(r'\[(J|C|M|D|EB/OL|R)\]', entry), \
            f'[{num}] 缺文献类型标识([J]/[C]/[EB/OL] 等): {entry}'


def test_reference_placeholder_is_gone():
    """「待补」占位符必须已经填掉 —— 交上去空着就是事故。"""
    text = _proposal_text()
    assert '待补' not in text, '开题报告书里还留着「待补」占位符'


def test_no_unverified_reference_markers_remain():
    """「待核」标记必须清零。

    写文献表时凡是我没亲眼核到原始页面的字段,一律标「(页码待核)」,
    提醒自己定稿前复核 —— 交上去时这些标记必须全部消掉。
    """
    text = _proposal_text()
    assert '待核' not in text, '参考文献里还有未核实的字段标记「待核」'
